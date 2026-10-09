from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

OUT=Path(__file__).resolve().parent
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.7)
sec.left_margin=sec.right_margin=Inches(.8)
sec.header_distance=sec.footer_distance=Inches(.3)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(7)
normal.paragraph_format.line_spacing=1.08
for st in doc.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for name,size in [('Title',27),('Subtitle',13),('Heading 1',19),('Heading 2',13),('Heading 3',11)]:
    st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(7)
doc.styles['Title'].paragraph_format.space_before=Pt(0)
header=sec.header.paragraphs[0]
header.text='GARAJE INTELIGENTE UNO     |     MATERIAL DE INVESTIGACIÓN'
header.runs[0].font.size=Pt(8);header.runs[0].font.color.rgb=RGBColor(0,0,0)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Guía de apoyo • 7 de octubre de 2026 • ').font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def p(text='',style=None):return doc.add_paragraph(text,style)
def h(text,level=2):return doc.add_heading(text,level)
def page(title):doc.add_page_break();h(title,1)
def bullet(text):return p(text,'List Bullet')
def step(text):return p(text,'List Number')
def formula(text):
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    m=OxmlElement('m:oMath');r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);m.append(r);para._p.append(m)
def table(headers,rows,widths):
    t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for i,txt in enumerate(headers):t.rows[0].cells[i].text=txt
    for row in rows:
        cells=t.add_row().cells
        for i,txt in enumerate(row):cells[i].text=str(txt)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for i,cell in enumerate(row.cells):
            cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'334B5E' if ri==0 else ('F1F4F6' if ri%2==0 else 'FFFFFF'));pr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2)
                para.paragraph_format.line_spacing=1.0
                for r in para.runs:r.font.size=Pt(10);r.font.bold=ri==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
    p().paragraph_format.space_after=Pt(0)
    return t
def link(paragraph,label,url):
    rel=paragraph.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    x=OxmlElement('w:hyperlink');x.set(qn('r:id'),rel)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'24536B');pr.append(c)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'19');pr.append(size)
    u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);r.append(pr);tx=OxmlElement('w:t');tx.text=url;r.append(tx);x.append(r);paragraph._p.append(x)

p('Investigación escolar sobre un garaje inteligente con Arduino UNO','Title')
p('Fundamentos técnicos y guía experimental para estudiantes en Paraguay','Subtitle')
p('Este material ayuda a transformar la construcción de una maqueta en una investigación comprobable. Reúne antecedentes, fuentes técnicas y un procedimiento para medir la detección y la respuesta de la puerta. Los resultados deben surgir de los ensayos del equipo; una demostración que funciona una vez no establece su confiabilidad.')
h('Qué se estudiará')
p('La maqueta detecta un objeto frente al HC-SR04, acciona uno o dos servos SG90 y registra presencia con un sensor infrarrojo interior. Un Arduino UNO ejecuta la lógica; la computadora ofrece configuración y acceso desde el celular. En el experimento, el objeto representa un vehículo. El sistema no reconoce automóviles ni identifica personas autorizadas.')
h('Cómo utilizar esta guía')
bullet('Lean las fuentes y escriban el problema con sus propias palabras. Completen institución, localidad, curso, integrantes y docente orientador.')
bullet('Acuerden una pregunta, variables y criterios antes de medir. Conserven el mismo montaje y configuración durante cada serie.')
bullet('Registren todos los intentos, incluidos fallos y lecturas inválidas. Elaboren gráficos y conclusiones limitadas a lo observado.')
bullet('Adjunten fotografías del montaje, versiones del programa, planillas y un registro de las tareas de cada integrante.')
h('Alcance de la evidencia disponible')
p('El registro local del proyecto documenta compilación para UNO, 26 comprobaciones de la lógica y 19 pruebas del servidor. También describe una revisión de la interfaz y del servidor portable. Son evidencias de software. Aún deben comprobarse sensores, alimentación, movimiento, EEPROM y acceso desde un teléfono físico (documentación local, 2026).')
p('Las cantidades de ensayos y metas propuestas aquí son decisiones didácticas. No se presentan como reglamento de una feria ni como certificación de un portón real. Este documento apoya la investigación; no reemplaza el informe que redactarán los estudiantes.')
p('Institución y localidad: __________________________________________\nCurso y equipo: _________________________________________________\nDocente y periodo de trabajo: _____________________________________')

page('1 Contexto paraguayo y antecedentes')
h('Conectividad y funcionamiento sin Internet')
p('El INE (2025, p. 9) informa que en 2024 utilizó Internet el 81,6 % de la población de 10 años y más cubierta por la encuesta: 86,2 % en el área urbana y 73,7 % en la rural. El indicador refiere al uso en los tres meses anteriores a la entrevista. La publicación excluye Boquerón, Alto Paraguay, comunidades indígenas y viviendas colectivas (p. 10). No equivale al porcentaje de hogares con Wi-Fi ni demuestra disponibilidad en una escuela concreta.')
p('Como decisión de diseño, conviene que la maqueta conserve su automatismo sin Internet y que el teléfono sea una vía adicional de control. Esto debe comprobarse en la institución: registrar qué red existe, si PC y teléfono pueden comunicarse y qué sucede cuando se detiene el servidor. La autonomía del programa no implica autonomía energética: el UNO sigue necesitando alimentación USB.')
h('Antecedente técnico nacional')
p('Burgos Delvalle y Estigarribia Barreto (2020), vinculados a la Universidad Nacional de Caaguazú, presentan una maqueta doméstica con Arduino UNO, Bluetooth HC-06 y una aplicación con comandos de voz. Reportan dificultades por ruido y dependencia de datos para el reconocimiento de voz, con botones como alternativa. El artículo es pertinente por su integración de control electrónico, vivienda a escala y teléfono.')
p('En este garaje, el disparador principal es un sensor de distancia y el control remoto pasa por una computadora. El aporte escolar puede ser medir cuándo detecta correctamente y documentar su comportamiento ante fallos. El antecedente no prueba el rendimiento de nuestra puerta ni permite afirmar que sea el primer sistema de este tipo en Paraguay. La ficha editorial publica el artículo en 2020 dentro de un volumen rotulado 2019; se cita el año de publicación registrado.')
h('Antecedente educativo nacional')
p('El póster Pequeños Inventores, de Piensa E.A.S. (2026), alojado en el repositorio de CONACYT, describe formación y mentoría en robótica para jóvenes de Alto Paraná. Sirve como ejemplo del uso educativo de la robótica en el país. Es una descripción de programa, no un experimento que demuestre mejoras de aprendizaje ni un antecedente específico de portones.')
h('Observación local que falta realizar')
p('Describan el lugar y necesidad que motivan el proyecto: por ejemplo, estudiar un acceso automatizado en una maqueta escolar. Registren materiales disponibles, condiciones de luz y red y costos en guaraníes. No atribuyan a toda la población paraguaya las preferencias de un grupo pequeño, ni afirmen que la maqueta reduce robos, energía o accidentes sin un estudio que lo mida.')

page('2 Fundamentos de sensores y control')
h('Arduino y automatización')
p('El UNO R3 utiliza el ATmega328P, dispone de 14 pines digitales y 6 entradas analógicas y trabaja con un reloj de 16 MHz. Su memoria EEPROM permite conservar datos sin alimentación (Arduino, s. f.-c). En este proyecto ejecuta una secuencia de reglas; el nombre “inteligente” describe la respuesta automática a sensores, no un sistema de inteligencia artificial.')
h('Distancia por ultrasonido')
p('El HC-SR04 emite una señal acústica y mide el tiempo de regreso del eco. La distancia se obtiene considerando el recorrido de ida y vuelta. La hoja de ELECFREAKS (s. f.) especifica alimentación de 5 V, disparo de al menos 10 µs y alcance nominal de 2 a 400 cm. También indica la conversión aproximada del tiempo de eco en microsegundos a centímetros al dividirlo por 58.')
formula('d ≈ t / 58')
p('Aquí d se expresa en centímetros y t en microsegundos. La hoja menciona una exactitud que puede llegar a 3 mm; no debe copiarse como resultado logrado con un automóvil de juguete. Su tamaño, forma y orientación pueden cambiar el eco. El firmware actual entrega centímetros enteros y considera inválida la ausencia de eco. La resolución que muestra el programa tampoco equivale a exactitud.')
h('Presencia por reflexión infrarroja')
p('Un módulo reflectivo emite luz infrarroja y detecta parte de la luz que devuelve el objeto. En el ejemplo técnico de Keyestudio (s. f.), un potenciómetro ajusta el umbral y la salida es digital. Esta fuente explica el principio, pero el KS0051 tiene cuatro pines y no identifica necesariamente el módulo de las imágenes del proyecto. Hay que anotar la marca, el modelo y los rótulos reales antes de transferir especificaciones.')
p('En la maqueta se registra presencia o ausencia, no distancia en centímetros. Se comprobará el nivel activo LOW o HIGH y la respuesta al vehículo. Como el IR mira hacia el interior, puede permanecer activo con el automóvil estacionado; no vigila toda la zona bajo la puerta.')
h('Confirmación e histéresis')
p('El código exige que la detección exterior dure un tiempo antes de aceptarla. Además, utiliza un umbral diferente para confirmar que la zona quedó libre: con máximo de 25 cm y margen de 5 cm, requiere una lectura mayor que 30 cm. Esta separación reduce cambios de estado cerca del umbral. Son reglas del proyecto, verificables en control.h; su eficacia física debe medirse.')

page('3 Movimiento alimentación y memoria')
h('Qué ordena un servo')
p('La biblioteca Servo permite ordenar pulsos en microsegundos. Leer su consigna no informa la posición física alcanzada (Arduino, s. f.-b). Por eso la interfaz muestra una apertura estimada; una puerta trabada puede permanecer inmóvil aunque la pantalla llegue a 100 %. Los alumnos deben observar o filmar el movimiento real.')
p('Tower Pro (s. f.) publica para su SG90 Digital un par de bloqueo de 1,8 kgf·cm a 4,8 V. Se trata del modelo descrito por ese fabricante y de una condición de bloqueo, no de una carga continua garantizada. Los motores de la maqueta deben identificarse y probarse sin forzar topes. No se asume que todos los servos vendidos como SG90 tengan el mismo recorrido o desempeño.')
h('Peso y brazo de palanca')
p('La exigencia al motor depende tanto de la masa de la puerta como de la distancia de su centro de masa al eje. Una estimación del momento gravitatorio máximo se obtiene al multiplicar masa, gravedad y brazo perpendicular; no incluye fricción, aceleración ni pérdidas del mecanismo.')
formula('τ ≈ m · g · r')
p('Ejemplo de cálculo, no medición: una hoja de 0,10 kg con centro de masa a 0,08 m del eje puede exigir aproximadamente 0,078 N·m, equivalentes a 0,80 kgf·cm. Esta estimación no autoriza trabajar al límite del servo. Un contrapeso o una hoja más liviana pueden reducir la exigencia.')
p('Con dos servos enfrentados se calibran extremos independientes y movimientos opuestos. Ambos actúan sobre el mismo eje: una diferencia de montaje o recorrido puede hacer que se opongan. Antes de acoplarlos se verifican por separado extremos y posiciones intermedias. No se concluye que dos motores dupliquen la capacidad útil.')
h('Alimentación del montaje')
p('Se acordó usar cuatro AA de 1,5 V para los servos y USB para el UNO. El montaje previsto incorpora regulación externa a 5 V y GND común. El positivo del portapilas no se conecta al pin 5V del UNO. Deben registrarse tensión en reposo y durante movimiento; el valor nominal de las pilas no demuestra estabilidad bajo carga. No se ensaya con la red eléctrica domiciliaria.')
h('Guardar no es medir')
p('EEPROM.put utiliza actualización de las celdas que cambian (Arduino, s. f.-a). En este proyecto, Aplicar modifica RAM y Guardar conserva configuración en EEPROM. La prueba consiste en guardar, reiniciar y volver a leer los valores. Hacerlo en demo solo verifica la simulación. La pausa se conserva por separado y bloquea el arranque automático hasta reanudar.')

page('4 Descripción verificable del sistema')
p('La siguiente descripción corresponde a la versión GARAGE-1.0.0 del código local. Copien en su informe los valores realmente utilizados; no presenten los valores iniciales como calibración de su maqueta.')
table(['Parte','Función y configuración inicial'],[
('HC-SR04 exterior','TRIG D4 y ECHO D2; rango de apertura 3–25 cm; confirmación de 300 ms.'),
('IR interior','OUT D7; nivel activo inicial LOW; registra presencia dentro.'),
('Servo 1 y servo 2','Señales D9 y D10. Modo inicial de un servo; extremos por calibrar.'),
('Exterior libre','Distancia válida mayor que 30 cm durante 1500 ms.'),
('Tiempos de espera','5 s si se registró presencia interior; 10 s si no hubo entrada. Solo cuenta con exterior libre.'),
('Recorrido','2500 ms de progresión de la orden; no equivale a tiempo físico garantizado.'),
('Arranque','Automatismo al encender deshabilitado inicialmente; requiere calibración y habilitación explícita.')],[1.65,5.25])
h('Secuencia observable')
p('Una vez inicializada y cerrada, la puerta abre al confirmarse un objeto exterior en el rango. Durante apertura y espera puede registrar presencia interior. Con exterior libre confirmado, espera el tiempo correspondiente y cierra. Si durante el cierre el exterior deja de estar libre o la lectura es inválida, ordena reabrir. Ante falta de eco permanece inhibido el cierre; es necesario un fondo que produzca una lectura válida fuera del rango de detección.')
h('Interfaz y panel')
p('El panel de Windows administra servidor, PIN, QR y ngrok. La interfaz web muestra mediciones y envía órdenes. El UNO se comunica por USB con la PC; no tiene conexión Wi-Fi propia en este montaje. La red local puede funcionar sin Internet; el acceso por ngrok sí depende de Internet y de la computadora encendida. El ciclo autónomo permanece en el UNO mientras esté alimentado.')
h('Límites que deben constar en el informe')
p('Pausar bloquea el automatismo y mantiene la última orden; no corta la energía. Mantenimiento retira las señales, por lo que hay que sostener la puerta. Tras reiniciar se desconoce su posición; el primer movimiento puede ser brusco. El IR interior no es una barrera de protección del umbral. El PIN controla el acceso a la interfaz, pero la apertura automática por proximidad no autentica al vehículo.')

page('5 Pregunta objetivos e hipótesis')
h('Problema propuesto')
p('Una maqueta puede abrir ante un objeto cercano y aun así fallar en otras posiciones. Se necesita determinar en qué condiciones detecta y abre de forma repetible, y comprobar si el cierre respeta las condiciones programadas. El problema es técnico y puede estudiarse con una regla, observación y un registro de ensayos.')
h('Pregunta principal')
p('¿Con qué frecuencia y en cuánto tiempo inicia la apertura esta maqueta al colocar el mismo vehículo a 5, 15, 20 y 35 cm del sensor, manteniendo fijos la orientación, la alimentación y el rango programado de 3 a 25 cm?')
h('Objetivo general')
p('Evaluar la detección y la respuesta de apertura de la maqueta bajo condiciones controladas y documentar sus límites de funcionamiento.')
h('Objetivos específicos')
bullet('Comparar la proporción de aperturas a las cuatro distancias de referencia.')
bullet('Medir el tiempo desde la colocación estable del vehículo hasta el inicio físico del movimiento.')
bullet('Comprobar el cierre tras entrada, la aproximación abandonada y la reacción frente a lecturas inválidas.')
bullet('Verificar conservación de ajustes y funcionamiento sin interfaz, con alimentación mantenida.')
h('Hipótesis operativa para evaluar')
p('Con el vehículo orientado de frente y el sistema previamente calibrado, se espera inicio de apertura en un máximo de 2 s en al menos 9 de 10 intentos a cada distancia interior al rango (5, 15 y 20 cm), y ninguna apertura durante 5 s en los 10 intentos a 35 cm. Son metas propuestas para esta investigación, no prestaciones del fabricante. Se fijan antes de observar los datos.')
table(['Tipo de variable','Definición en este ensayo'],[
('Independiente','Distancia real de referencia: 5, 15, 20 o 35 cm.'),
('Dependientes','Apertura dentro del plazo, tiempo de respuesta y proporción de lecturas inválidas.'),
('Controladas','Vehículo, orientación, altura del sensor, masa de puerta, número de servos, configuración, fondo e iluminación.'),
('Registrar como condición','Fecha, tensión de alimentación, estado de pilas y cambios involuntarios del montaje.')],[1.6,5.3])
p('La unidad de observación es un intento reiniciado. Diez lecturas consecutivas del sensor durante una sola aproximación no son diez intentos independientes. Las repeticiones se realizan sobre una misma maqueta: no representan a todos los sensores o viviendas.')

page('6 Procedimiento del experimento principal')
p('Tipo de estudio: investigación aplicada con ensayos controlados de un prototipo. Se propone un factor principal, cuatro distancias y diez repeticiones por distancia: 40 intentos. El orden aleatorio evita hacer siempre una condición con pilas nuevas y otra al final; este criterio se apoya en NIST/SEMATECH (s. f.).')
h('Preparación')
step('Calibren sin carga y luego verifiquen la puerta bajo supervisión. Utilicen la guía MONTAJE_Y_CALIBRACION.md. Suspendan ante bloqueo, calentamiento o esfuerzo entre servos.')
step('Fijen el sensor y marquen distancias desde un plano de referencia definido en su cara frontal hasta la parte frontal del vehículo. Registren la resolución de la regla y la orientación usada.')
step('Configuren 3–25 cm, confirmación de 300 ms y margen de 5 cm. Mantengan los demás valores anotados. Inicialicen la puerta y esperen cierre físico completo; ese movimiento inicial no cuenta como ensayo.')
step('Preparen 40 tarjetas: diez por cada distancia. Mézclenlas y extraigan una por intento, sin reposición. Si una variación de alimentación obliga a interrumpir, registren el bloque y no oculten el cambio de pilas.')
h('Medición de cada intento')
step('Comiencen con puerta cerrada, acceso libre y eco válido de fondo. Mantengan el mismo estado del IR. Coloquen el vehículo desde un lado en la marca sorteada, evitando cruzar frente al sensor con la mano.')
step('Definan t0 como el instante en que el vehículo queda estable en la marca. Filmen simultáneamente vehículo y puerta. Mantengan la posición por 5 s; anoten si inicia apertura dentro de 2 s y si hubo alguna apertura hasta los 5 s.')
step('Registren una lectura de distancia alrededor de 0,5 s después de t0 y su validez. Si no hay eco válido, escriban “sin eco”; no lo conviertan en cero ni descarten el intento. Guarden el video o evidencia con el mismo identificador de la fila.')
step('Retiren el vehículo y esperen a que concluya el ciclo. Verifiquen cierre físico antes del próximo intento. Anoten cualquier intervención, reinicio o evento inesperado.')
h('Cómo medir tiempo sin confundir pantallas')
p('El tiempo físico de respuesta es la diferencia entre el primer movimiento visible de la puerta y t0. Con video, dividan los fotogramas transcurridos por los fotogramas por segundo reales de la grabación. Con cronómetro manual, informen el error de reacción del observador. La actualización de la interfaz incorpora demora de comunicación y no sustituye esta medición.')
p('Al acabar, revisen que existan 10 intentos por condición. No repitan un fallo hasta conseguir éxito para reemplazarlo. Si mejoran el montaje, conserven la serie original y comiencen otra serie identificada con la nueva configuración.')

page('7 Pruebas funcionales y extensiones')
p('Estas comprobaciones complementan el experimento de distancia. Se propone repetir cinco veces cada escenario y registrar resultado físico y telemetría por separado. Los resultados todavía están pendientes; ninguna casilla debe marcarse por haber pasado una prueba de software.')
table(['Escenario','Procedimiento y resultado esperado'],[
('Entrada normal','Detectar afuera, abrir, colocar el vehículo ante IR y liberar exterior. Debe iniciar cierre después de la espera configurada con exterior libre.'),
('Aproximación abandonada','Detectar afuera y retirar sin activar IR. Debe cerrar tras la espera sin entrada y la confirmación de exterior libre.'),
('Exterior ocupado','Con puerta abierta, mantener objeto detectado afuera más tiempo que la espera. El cierre debe permanecer inhibido.'),
('Obstáculo durante cierre','Con un objeto liviano fuera del recorrido mecánico, provocar detección exterior mientras cierra. Debe ordenar reapertura. No usar manos bajo la puerta.'),
('Falta de eco','Orientar temporalmente el sensor hacia una zona que no devuelva eco, sin introducir cortocircuitos. Debe inhibir cierre o reabrir si ya estaba cerrando.'),
('Memoria y autonomía','Guardar ajustes y leer tras reiniciar. Probar pausa persistida y después funcionamiento con servidor detenido, manteniendo alimentación y configuración apropiadas.')],[1.65,5.25])
p('Para temporización de cierre, registren cuándo la puerta queda abierta y cuándo se confirma exterior libre. Si el exterior ya estaba libre al completar apertura, la espera parte de la llegada al estado abierto; si no, comienza cuando se libera. Comparen inicio físico del cierre con la configuración y describan la incertidumbre de su medición.')
h('Extensión opcional sobre orientación')
p('A 15 cm y con el mismo vehículo, prueben orientaciones de 0°, 20° y 40° respecto de la posición frontal, diez veces cada una. Definan el punto de giro y conserven la distancia de referencia. Esta serie puede mostrar sensibilidad a la geometría; no debe mezclarse con la serie principal como si solo cambiara distancia.')
h('Extensión opcional sobre uno o dos servos')
p('Comparen una misma hoja y masa únicamente después de calibrar ambos modos. Midan tiempo físico y tensión bajo carga. Con un solo servo, desacoplen mecánicamente el segundo: dejarlo unido altera la resistencia. Nunca añadan peso hasta bloquear el eje. Registren masas y cambios de mecanismo; una diferencia de resultados no demuestra por sí sola una duplicación del par.')

page('8 Instrumentos de registro')
h('Ficha de la sesión')
p('Fecha y responsables: ____________________________________________\nVersión del código y archivo de configuración: _________________________\nModo de servos y masa de puerta: ___________________________________\nVehículo y orientación: ____________________________________________\nRegla o instrumento y resolución: ___________________________________\nAlimentación y tensión antes y después: ______________________________\nFondo ultrasónico e iluminación: ____________________________________')
p('Tabla 1. Registro por intento. Copiar esta página hasta completar la serie. En “apertura” anotar si/no dentro de 2 s y cualquier apertura tardía hasta 5 s en observaciones. Tiempo: desde t0 al primer movimiento físico; si no abre, escribir “no observado en 5 s”.')
table(['ID','Ref. cm','Sensor cm o sin eco','Abrió en 2 s','Tiempo s','Observaciones y evidencia'],[['','','','','',''] for _ in range(7)],[.45,.65,1.1,.8,.85,3.05])
p('La lectura del sensor se toma aproximadamente a 0,5 s de t0. Conserven la condición inválida como dato. El detalle de validez puede anotarse en observaciones. Una fotografía aislada no demuestra que se respetaron los tiempos del ciclo.')
h('Ficha breve de un fallo')
p('ID y condición: __________________________________________________\nQué se esperaba y qué ocurrió: _____________________________________\nEstado mostrado y movimiento observado: ___________________________\nPosible causa y evidencia disponible: _________________________________\nCambio realizado y nueva serie de prueba: ____________________________')
p('Bitácora del equipo: anotar fecha, tarea, responsable, decisión, archivo o fotografía y pendiente. Distinguir quién diseñó, programó, adaptó código, construyó y realizó ensayos. Reconocer el uso de bibliotecas y asistencia de IA; no atribuir a los estudiantes pruebas que no ejecutaron.')

page('9 Análisis conclusiones y pertinencia local')
h('Cálculos que pueden realizar')
formula('Aperturas dentro de plazo (%) = 100 · éxitos / intentos')
formula('Error absoluto (cm) = |distancia leída − distancia de referencia|')
p('Calculen apertura dentro de 2 s para cada distancia con denominador 10. A 35 cm, registren por separado cualquier apertura en la ventana completa de 5 s como apertura no esperada. Para tiempos, informen mediana o promedio y mínimo–máximo de los intentos que sí abrieron, indicando cuántos fueron; no asignen tiempo cero a los fallos.')
p('Calculen el error de distancia solo con lecturas válidas y publiquen también cuántas fueron inválidas. La lectura entera en centímetros y la regla limitan la resolución del ensayo. Si todos los intentos pasan, escriban “10 de 10 en estas condiciones”, no “100 % confiable en cualquier situación”.')
h('Presentación de resultados')
p('Un gráfico de barras puede comparar aperturas por distancia; una tabla resume n, éxitos, lecturas inválidas y tiempos. Otro gráfico compara distancia leída y distancia de referencia. No mezclen las pruebas simuladas con los ensayos físicos. Antes y después de una mejora deben identificarse como series diferentes.')
h('Redacción de la conclusión')
p('Modelo para completar: “En [n] intentos por condición, obtuvimos [resultados]. La hipótesis [recibió apoyo / no recibió apoyo] bajo [condiciones]. Los fallos se concentraron en [observaciones]. La limitación principal fue [evidencia]. Proponemos comprobar [mejora] en una nueva serie”. No es necesario que la hipótesis se cumpla para que la investigación sea válida.')
h('Costos y uso en Paraguay')
p('Soliciten cotizaciones locales fechadas en guaraníes. Separen materiales reutilizados de compras nuevas e incluyan UNO, sensores, servos, regulador, portapilas, pilas, cables y estructura. Registren por separado computadora y teléfono aunque ya estén disponibles. No comparen el costo de una maqueta con el de un portón instalado como si ofrecieran las mismas funciones.')
p('Una ficha de cotización debe tener componente, modelo, cantidad, precio unitario, subtotal, proveedor, fecha y costo de envío. Para argumentar ahorro necesitan una alternativa comparable y el mismo alcance. Registren también consumo de pilas; no supongan beneficios ambientales por usar automatización.')
h('Informe escolar sugerido')
p('Portada; resumen redactado al final; problema y objetivos; antecedentes; fundamentos; metodología; montaje y configuración; resultados; discusión y limitaciones; conclusiones; referencias y anexos. El docente puede adaptar esta estructura. No se exige una encuesta social para responder una pregunta técnica: los ensayos son la evidencia principal.')

page('10 Referencias técnicas y bibliográficas')
p('Referencias en formato orientativo APA 7. Fuentes consultadas el 7 de octubre de 2026. Las letras de Arduino distinguen documentos sin fecha. Los enlaces permiten revisar el contenido original; las notas indican qué aporta cada fuente.')
refs=[
('Arduino. (s. f.-a). EEPROM Library V2.0 for Arduino [Documentación de software]. GitHub.', 'https://github.com/arduino/ArduinoCore-avr/blob/master/libraries/EEPROM/README.md','Sustenta la diferencia entre leer, escribir y actualizar EEPROM. Consultar put y update.'),
('Arduino. (s. f.-b). Servo library [Documentación de software]. GitHub.', 'https://github.com/arduino-libraries/Servo/blob/master/docs/api.md','Consultar writeMicroseconds y read: la consigna no es realimentación de posición física.'),
('Arduino. (s. f.-c). UNO R3 [Documentación de hardware].','https://docs.arduino.cc/hardware/uno-rev3/','Identifica la placa y sus recursos. No confundir UNO R3 con UNO R4.'),
('Burgos Delvalle, D., & Estigarribia Barreto, H. R. (2020). Domótica de bajo coste controlada por comandos de voz. Tecnología, Diseño e Innovación, 5(1).','https://www.unae.edu.py/ojs/index.php/facat/article/view/155','Antecedente paraguayo. La ficha fecha publicación el 14 de enero de 2020 y el número se rotula 2019. No se infieren precios actuales.'),
('ELECFREAKS. (s. f.). Ultrasonic ranging module HC-SR04 [Hoja de datos].','https://www.elecfreaks.com/download/HC-SR04.pdf','Fundamento de tiempo de vuelo, interfaz y condiciones nominales. La exactitud del fabricante no reemplaza la medición escolar.'),
]
for txt,url,note in refs:
    a=p(txt);a.paragraph_format.left_indent=Inches(.22);a.paragraph_format.first_line_indent=Inches(-.22)
    link(p(),'Abrir fuente original',url);p('Uso en la investigación: '+note)

page('11 Fuentes de contexto y trazabilidad')
refs2=[
('Instituto Nacional de Estadística. (2025). Tecnología de la información y comunicación en el Paraguay 2024.','https://www.ine.gov.py/Publicaciones/Biblioteca/documento/280/Tics%202024_INE.pdf','Datos de 2024, publicación de junio de 2025. Revisar población cubierta y páginas 9–10 antes de citar porcentajes.'),
('Keyestudio. (s. f.). KS0051 keyestudio infrared obstacle avoidance sensor [Documentación técnica].','https://wiki.keyestudio.com/Ks0051_keyestudio_Infrared_Obstacle_Avoidance_Sensor','Referencia del principio reflectivo. Es un módulo de cuatro pines; no usar su pinout como identificación del sensor de esta maqueta.'),
('NIST/SEMATECH. (s. f.). Completely randomized designs. En e-Handbook of Statistical Methods (sección 5.3.3.1).','https://www.itl.nist.gov/div898/handbook/pri/section3/pri331.htm','Apoya la planificación de niveles, repeticiones y orden aleatorio. Las 40 pruebas propuestas son una adaptación didáctica propia.'),
('Piensa E.A.S. (2026, 13 de abril). Pequeños Inventores [Póster]. Repositorio CONACYT.','https://repositorio.conacyt.gov.py/handle/20.500.14066/4798','Antecedente educativo en Alto Paraná; no representa un resultado experimental del garaje.'),
('Tower Pro. (s. f.). SG90 Digital [Ficha de producto].','https://towerpro.com.tw/product/sg90-7/','Distinguir datos del modelo original, par de bloqueo y capacidad útil del mecanismo construido.'),
]
for txt,url,note in refs2:
    a=p(txt);a.paragraph_format.left_indent=Inches(.22);a.paragraph_format.first_line_indent=Inches(-.22)
    link(p(),'Abrir fuente original',url);p('Uso en la investigación: '+note)
h('Documentos primarios del propio proyecto')
p('Garaje Inteligente UNO. (2026). Firmware GARAGE-1.0.0 y documentación de validación [Código y documentación local no publicados]. Archivos: firmware/GarajeInteligente/control.h, GarajeInteligente.ino y docs/VALIDACION.md. El equipo debe completar la autoría y la versión efectivamente usada en su informe.')
p('Conserven una copia del código ensayado, la configuración, datos crudos y evidencias. Citen solo fuentes que hayan leído. Una respuesta de IA puede ayudar a organizar el trabajo, pero no sustituye la ficha técnica ni los resultados experimentales.')

doc.core_properties.title='Investigación escolar sobre un garaje inteligente con Arduino UNO'
doc.core_properties.subject='Fundamentos y protocolo experimental en contexto paraguayo'
doc.core_properties.author='Material de apoyo para el equipo escolar'
path=OUT/'Guia_investigacion_garaje_Paraguay.docx'
doc.save(path)
print(path)
