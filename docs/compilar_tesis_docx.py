import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

def create_element(name):
    return OxmlElement(name)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, color="D0D5DD", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for b_name in ['top', 'left', 'bottom', 'right']:
        b = OxmlElement(f'w:{b_name}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), sz)
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), color)
        borders.append(b)
    tcPr.append(borders)

def clear_cell_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for b_name in ['top', 'left', 'bottom', 'right']:
        b = OxmlElement(f'w:{b_name}')
        b.set(qn('w:val'), 'none')
        borders.append(b)
    tcPr.append(borders)

def set_cell_shading(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)

def add_page_number(paragraph):
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    paragraph._p.append(fldSimple)

def parse_inline(paragraph, text, base_font_size=11, is_code=False, is_quote=False):
    if is_code:
        r = paragraph.add_run(text)
        r.font.name = 'Consolas'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(30, 41, 59)
        return

    # Pattern for bold, italic, inline code, latex math
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\$\$.*?\$\$|\$.*?\$)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            r = paragraph.add_run(token[2:-2])
            r.font.name = 'Times New Roman'
            r.font.size = Pt(base_font_size)
            r.font.bold = True
            if is_quote:
                r.font.italic = True
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            r = paragraph.add_run(token[1:-1])
            r.font.name = 'Times New Roman'
            r.font.size = Pt(base_font_size)
            r.font.italic = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            r = paragraph.add_run(token[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(base_font_size - 1)
            r.font.color.rgb = RGBColor(180, 40, 40)
        elif (token.startswith('$$') and token.endswith('$$') and len(token) >= 4) or (token.startswith('$') and token.endswith('$') and len(token) >= 2):
            clean_math = token.strip('$')
            # Clean up latex commands for clean Word presentation
            clean_math = clean_math.replace(r'\mathbf', '').replace(r'\text', '').replace(r'\rm', '')
            clean_math = clean_math.replace(r'\approx', '≈').replace(r'\le', '≤').replace(r'\ge', '≥')
            clean_math = clean_math.replace(r'\times', '×').replace(r'\cdot', '·').replace(r'\mu', 'µ')
            clean_math = clean_math.replace(r'\Delta', 'Δ').replace(r'\tau', 'τ').replace(r'\perp', '⊥')
            clean_math = clean_math.replace(r'\implies', '⇒').replace(r'\underline', '').replace(r'\hspace', '')
            clean_math = clean_math.replace(r'\vspace', '').replace(r'\quad', ' ')
            clean_math = re.sub(r'\\frac\{([^{}]+)\}\{([^{}]+)\}', r'(\1 / \2)', clean_math)
            clean_math = re.sub(r'\{([^{}]+)\}', r'\1', clean_math)
            clean_math = re.sub(r'\\([a-zA-Z]+)', r'\1', clean_math)
            r = paragraph.add_run(clean_math)
            r.font.name = 'Cambria Math'
            r.font.size = Pt(base_font_size)
            r.font.italic = True
            r.font.color.rgb = RGBColor(20, 40, 80)
        else:
            r = paragraph.add_run(token)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(base_font_size)
            if is_quote:
                r.font.italic = True

def compile_thesis(md_path, docx_path):
    doc = Document()
    
    # Configure margins: 1 inch (2.54 cm) around
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)
        sec.header_distance = Inches(0.5)
        sec.footer_distance = Inches(0.5)
        
        # Header setup
        header = sec.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("GARAJE INTELIGENTE UNO  ·  SERCAP 2026  ·  BTI CELE")
        hrun.font.name = 'Times New Roman'
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(120, 130, 140)
        
        # Footer setup
        footer = sec.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun1 = fp.add_run("Página ")
        frun1.font.name = 'Times New Roman'
        frun1.font.size = Pt(9)
        frun1.font.color.rgb = RGBColor(100, 110, 120)
        add_page_number(fp)
        
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []
    
    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        
        # Check code fence
        if line.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lines = []
            else:
                in_code_block = False
                # Render code block
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                set_cell_shading(cell, "F8F9FA")
                set_cell_borders(cell, "E2E8F0", "6")
                set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                cp = cell.paragraphs[0]
                cp.paragraph_format.line_spacing = 1.05
                cp.paragraph_format.space_before = Pt(2)
                cp.paragraph_format.space_after = Pt(2)
                for cl in code_lines:
                    parse_inline(cp, cl + '\n', is_code=True)
                # spacing after table
                sp = doc.add_paragraph()
                sp.paragraph_format.space_after = Pt(4)
                sp.paragraph_format.line_spacing = 1.0
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue
            
        # Table detection
        if line.strip().startswith('|') and line.strip().endswith('|'):
            table_lines.append(line.strip())
            i += 1
            while i < len(lines) and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            # Process table
            parsed_rows = []
            for tl in table_lines:
                raw_cells = [c.strip() for c in tl.split('|')[1:-1]]
                if all(re.match(r'^:?-+:?$', c) for c in raw_cells if c):
                    continue # separator line
                parsed_rows.append(raw_cells)
                
            if parsed_rows:
                num_cols = max(len(r) for r in parsed_rows)
                for r in parsed_rows:
                    while len(r) < num_cols:
                        r.append('')
                
                # Check if it's a signature table
                is_signature = any('Firma del' in c or '______' in c for r in parsed_rows for c in r)
                
                table = doc.add_table(rows=len(parsed_rows), cols=num_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = True
                
                for r_idx, row_data in enumerate(parsed_rows):
                    row = table.rows[r_idx]
                    trPr = row._tr.get_or_add_trPr()
                    cantSplit = OxmlElement('w:cantSplit')
                    trPr.append(cantSplit)
                    
                    if r_idx == 0 and not is_signature:
                        tblHeader = OxmlElement('w:tblHeader')
                        trPr.append(tblHeader)
                        
                    for c_idx, cell_value in enumerate(row_data):
                        cell = row.cells[c_idx]
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        
                        if is_signature:
                            set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
                            clear_cell_borders(cell)
                            set_cell_shading(cell, "FFFFFF")
                            p = cell.paragraphs[0]
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            p.paragraph_format.line_spacing = 1.15
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            parse_inline(p, cell_value, base_font_size=10.5)
                        else:
                            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                            set_cell_borders(cell, "CBD5E1", "4")
                            
                            if r_idx == 0:
                                set_cell_shading(cell, "1E3A5F") # Navy blue header
                            elif r_idx % 2 == 1:
                                set_cell_shading(cell, "FFFFFF")
                            else:
                                set_cell_shading(cell, "F1F5F9") # Soft slate alternating
                                
                            p = cell.paragraphs[0]
                            p.paragraph_format.line_spacing = 1.15
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            
                            if r_idx == 0:
                                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                r = p.add_run(cell_value)
                                r.font.name = 'Times New Roman'
                                r.font.size = Pt(10)
                                r.font.bold = True
                                r.font.color.rgb = RGBColor(255, 255, 255)
                            else:
                                parse_inline(p, cell_value, base_font_size=9.5)
                
                # Spacing after table
                sp = doc.add_paragraph()
                sp.paragraph_format.space_after = Pt(6)
                sp.paragraph_format.line_spacing = 1.0
                
            table_lines = []
            continue

        # Check page breaks
        if line.strip() in ['\\newpage', '---']:
            if line.strip() == '\\newpage':
                doc.add_page_break()
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                pBdr = OxmlElement('w:pBdr')
                bottom = OxmlElement('w:bottom')
                bottom.set(qn('w:val'), 'single')
                bottom.set(qn('w:sz'), '6')
                bottom.set(qn('w:space'), '1')
                bottom.set(qn('w:color'), 'CBD5E1')
                pBdr.append(bottom)
                p._p.get_or_add_pPr().append(pBdr)
            i += 1
            continue

        # Headings
        if line.startswith('# '):
            h_text = line[2:].strip()
            # If major chapter, add page break
            if any(h_text.startswith(prefix) for prefix in ['CAPÍTULO', 'BLOQUE', 'REFERENCIAS', 'ANEXOS']):
                doc.add_page_break()
                
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(16)
            r.font.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)
            i += 1
            continue

        if line.startswith('## '):
            h_text = line[3:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(13.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(30, 41, 59)
            i += 1
            continue

        if line.startswith('### '):
            h_text = line[4:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = RGBColor(51, 65, 85)
            i += 1
            continue

        if line.startswith('#### '):
            h_text = line[5:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.italic = True
            r.font.color.rgb = RGBColor(71, 85, 105)
            i += 1
            continue

        # Blockquotes
        if line.startswith('> '):
            quote_text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.right_indent = Inches(0.5)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.3
            parse_inline(p, quote_text, base_font_size=10.5, is_quote=True)
            i += 1
            continue

        # Bullet list
        if line.strip().startswith('* ') or line.strip().startswith('- '):
            bullet_text = line.strip()[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.25
            parse_inline(p, bullet_text, base_font_size=11)
            i += 1
            continue

        # Checkbox item
        if line.strip().startswith('* [x] ') or line.strip().startswith('- [x] ') or line.strip().startswith('* [ ] ') or line.strip().startswith('- [ ] '):
            is_checked = '[x]' in line
            box_symbol = "☑ " if is_checked else "☐ "
            item_text = line.strip()[6:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.25
            r_box = p.add_run(box_symbol)
            r_box.font.name = 'Arial'
            r_box.font.size = Pt(11)
            r_box.font.bold = True
            parse_inline(p, item_text, base_font_size=11)
            i += 1
            continue

        # Numbered list
        num_match = re.match(r'^(\d+)\.\s+(.*)$', line.strip())
        if num_match:
            item_text = num_match.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.25
            parse_inline(p, item_text, base_font_size=11)
            i += 1
            continue

        # Math formulas on own line
        if line.strip().startswith('$$') and line.strip().endswith('$$') and len(line.strip()) >= 4:
            math_text = line.strip()[2:-2].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            parse_inline(p, f"$${math_text}$$", base_font_size=12)
            i += 1
            continue

        # Regular paragraph
        if line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.35
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            parse_inline(p, line.strip(), base_font_size=11)
        else:
            # blank line, ignore or leave small gap
            pass
            
        i += 1

    try:
        doc.save(docx_path)
        print(f"Documento DOCX compilado exitosamente en: {docx_path}")
    except PermissionError:
        alt_path = docx_path.with_name(docx_path.stem + "_Corregida.docx")
        doc.save(alt_path)
        print(f"El archivo original está abierto en Word. Se guardó la versión corregida en: {alt_path}")

if __name__ == '__main__':
    base_dir = Path(__file__).resolve().parent
    md_file = base_dir / 'Tesis_Garaje_Inteligente_CELE_2026.md'
    docx_file = base_dir / 'Tesis_Garaje_Inteligente_CELE_2026.docx'
    compile_thesis(md_file, docx_file)
